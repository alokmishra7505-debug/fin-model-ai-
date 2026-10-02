from dataclasses import replace
from io import BytesIO
from openpyxl import load_workbook
import pytest
from finmodel_agents.orchestrator import run_model
from finmodel_agents.export import excel_workbook
from finmodel_agents.utils import Units
from excel_evaluator import Evaluator


def test_all_excel_formulas_match_python_results(sample_data):
    b=run_model(sample_data)
    raw=excel_workbook(b,Units("USD"))
    formulas=load_workbook(BytesIO(raw),data_only=False)
    expected=load_workbook(BytesIO(raw),data_only=True)
    evaluator=Evaluator(formulas)
    for ws in formulas:
        for row in ws:
            for cell in row:
                if cell.data_type=="f":
                    actual=evaluator.cell(ws.title,cell.coordinate)
                    target=expected[ws.title][cell.coordinate].value
                    if isinstance(target,(int,float)):
                        assert actual==pytest.approx(target,abs=1e-7),f"{ws.title}!{cell.coordinate}"
                    else:
                        assert actual==target


@pytest.mark.parametrize("currency,selection",[("USD","Million"),("INR","Crore"),("USD","Billion")])
def test_excel_assumption_change_recalculates_forecast_and_dcf(sample_data,currency,selection):
    sample_data.info["currency"]=currency
    sample_data.info["financialCurrency"]=currency
    b=run_model(sample_data)
    units=Units(currency,selection)
    wb=load_workbook(BytesIO(excel_workbook(b,units)),data_only=False)
    wb["Assumptions"]["B6"]=.15
    wb["Assumptions"]["B13"]=.12  # Capex / revenue; meaningful cash and debt effects.
    evaluator=Evaluator(wb)
    expected=run_model(sample_data,replace(b.assumptions,growth=.15,capex_pct=.12))
    f=expected.forecasts["Base"]
    for name,frame in (("Forecast IS",f.income),("Forecast BS",f.balance),("Forecast CF",f.cashflow)):
        for r,(metric,values) in enumerate(frame.iterrows(),6):
            for c,value in enumerate(values,2):
                target=value if metric=="EPS" else value/units.scale
                assert evaluator.cell(name,wb[name].cell(r,c).coordinate)==pytest.approx(target,abs=1e-6)
    assert evaluator.cell("DCF","B7")==pytest.approx(expected.valuations["Base"].values["Implied Price"])
