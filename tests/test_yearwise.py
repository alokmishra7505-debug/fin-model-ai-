from dataclasses import replace
from io import BytesIO
import pytest
from openpyxl import load_workbook
from finmodel_agents.orchestrator import run_model
from finmodel_agents.export import excel_workbook
from finmodel_agents.utils import Units
from excel_evaluator import Evaluator

@pytest.mark.parametrize('case',['Base','Bear','Bull'])
def test_annual_drivers_selected_export_recalculate(sample_data,case):
    initial=run_model(sample_data)
    a=initial.assumptions
    annual=tuple(replace(a,growth=.04+n*.02,capex_pct=.03+n*.01,tax_rate=.2+n*.02) for n in range(5))
    bundle=run_model(sample_data,a,yearly=annual)
    raw=excel_workbook(bundle,Units('USD'),scenario=case)
    wb=load_workbook(BytesIO(raw),data_only=False)
    cached=load_workbook(BytesIO(raw),data_only=True)
    evaluator=Evaluator(wb)
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type=='f':
                    target=cached[ws.title][cell.coordinate].value
                    actual=evaluator.cell(ws.title,cell.coordinate)
                    if isinstance(target,(float,int)):
                        assert actual==pytest.approx(target,abs=1e-6),(ws.title,cell.coordinate)
                    else:
                        assert actual==target
    assert f'Selected case: {case}' in cached['Cover']['A12'].value
    wb['Assumptions']['D6']=.19
    ev=Evaluator(wb)
    assert ev.cell('Forecast IS','C6')==pytest.approx(ev.cell('Forecast IS','B6')*1.19)
    assert ev.cell('Forecast BS','C21') is not None
