from dataclasses import replace
from io import BytesIO
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from finmodel_agents.utils import Units,safe_div
from finmodel_agents.normalizer import normalize
from finmodel_agents.forecasting import Assumptions, default_assumptions
from finmodel_agents.three_statement import build_forecast
from finmodel_agents.valuation import wacc,dcf_value,value_company,monte_carlo,reverse_dcf
from finmodel_agents.orchestrator import run_model
from finmodel_agents.export import excel_workbook,powerbi_files,SHEETS
from finmodel_agents.market_data import MarketDataAgent
from finmodel_agents.resolver import resolve


def test_units_and_per_share():
    u = Units("INR")
    assert u.money(149000000000)=="₹14,900.0 Cr"
    assert u.money(-1250000000)=="(₹125.0 Cr)"
    assert u.money(1225.4,True)=="₹1,225.40"
    assert u.money(0)=="–"
    assert u.money(np.nan)=="N/A"
    assert Units("USD").money(1e9)=="$1,000.0m"
    assert Units("INR","Billion").scale==1e9
    assert np.isnan(safe_div(1,0))


def test_wacc_and_dcf_math():
    a=Assumptions(risk_free=.04,beta=1,equity_premium=.06,cost_debt=.05,tax_rate=.25)
    assert wacc(a,800,200)==pytest.approx(.0875)
    v=dcf_value([100]*5,.10,0,50,10)
    assert v["Enterprise Value"]==pytest.approx(1000)
    assert v["Equity Value"]==pytest.approx(950)
    assert v["Implied Price"]==pytest.approx(95)
    with pytest.raises(ValueError):
        dcf_value([100],.03,.03)
    assert np.isnan(wacc(a,np.nan,100))


@pytest.mark.parametrize("years",[3,4,5])
def test_statements_reconcile(sample_data,years):
    bundle=run_model(sample_data,years=years)
    for f in bundle.forecasts.values():
        assert len(f.income.columns)==years
        assert f.balance.loc["Balance Check"].abs().max()<1e-5
        assert f.cashflow.loc["Cash Reconciliation"].abs().max()<1e-5
        assert (f.balance.loc["Cash"].values==f.cashflow.loc["Closing Cash"].values).all()
        assert f.income.iloc[:,0]["Revenue"]==pytest.approx(1.1e9*(1+f.assumptions.growth))
        assert f.balance.iloc[:,0]["Equity"]==pytest.approx(5e8+f.income.iloc[:,0]["Net Income"]-f.cashflow.iloc[:,0]["Dividends"])
    assert not (bundle.checks.Status=="FAIL").any()


def test_explicit_funding_and_loss_taxes(sample_data):
    hist=normalize(sample_data)
    a=Assumptions(ebitda_margin=.01,capex_pct=.5,da_pct=.05,min_cash_pct=.2)
    f=build_forecast(hist,sample_data.info,a)
    assert f.cashflow.loc["Funding Draw"].sum()>0
    assert f.balance.loc["Balance Check"].abs().max()<1e-5
    assert (f.income.loc["Taxes"]==0).all()
    assert (f.balance.loc["Cash"]>=f.income.loc["Revenue"]*.2-1e-5).all()
    assert (f.balance.loc["Net PPE"]>=0).all()


def test_missing_components_are_disclosed(sample_data):
    sample_data.balance=sample_data.balance.drop(index="Inventory")
    hist=normalize(sample_data)
    assert hist.balance.loc["Inventory"].isna().all()
    f=build_forecast(hist,sample_data.info,Assumptions())
    assert any("inventory was unavailable" in x for x in f.notes)
    assert f.balance.loc["Balance Check"].abs().max()<1e-5


def test_unbalanced_opening_is_not_plugged(sample_data):
    sample_data.balance.loc["Total Assets"]+=1e8
    b=run_model(sample_data)
    assert not b.forecasts
    assert any("does not balance" in x for x in b.notices)


def test_missing_required_anchor_pauses_forecast(sample_data):
    sample_data.balance=sample_data.balance.drop(index="Total Assets")
    assert not run_model(sample_data).forecasts


def test_banks_and_currency_mismatch(sample_data):
    sample_data.info["sector"]="Financial Services"
    assert not run_model(sample_data).forecasts
    sample_data.info["sector"]="Technology"
    sample_data.info["financialCurrency"]="INR"
    b=run_model(sample_data)
    assert not b.valuations["Base"].values
    assert "currencies differ" in b.valuations["Base"].message


def test_negative_terminal_fcf_is_unavailable(sample_data):
    b=run_model(sample_data,Assumptions(ebitda_margin=.05,capex_pct=.5))
    assert not b.valuations["Base"].values


def test_sensitivity_center_and_reverse_dcf(sample_data):
    b=run_model(sample_data)
    v=b.valuations["Base"]
    assert v.sensitivity.iloc[2,2]==pytest.approx(v.values["Implied Price"])
    implied,_=reverse_dcf(b.historical,sample_data,b.assumptions,5,v.values["Implied Price"])
    assert implied==pytest.approx(b.assumptions.growth,abs=1e-7)


def test_monte_carlo_reproducible(sample_data):
    b=run_model(sample_data)
    x,s=monte_carlo(b.forecasts["Base"],sample_data,1000,42)
    y,_=monte_carlo(b.forecasts["Base"],sample_data,1000,42)
    assert np.array_equal(x,y)
    assert len(x)>900
    assert s["P10"]<s["Median / P50"]<s["P90"]


def test_excel_values_formulas_and_powerbi(sample_data):
    b=run_model(sample_data)
    payload=excel_workbook(b,Units("USD"))
    wb=load_workbook(BytesIO(payload),data_only=False)
    cached=load_workbook(BytesIO(payload),data_only=True)
    assert wb.sheetnames==SHEETS
    assert wb["Forecast IS"]["B6"].value.startswith("=")
    assert cached["Forecast IS"]["B6"].value==pytest.approx(b.forecasts["Base"].income.iloc[0,0]/1e6)
    assert cached["DCF"]["B7"].value==pytest.approx(b.valuations["Base"].values["Implied Price"])
    for ws in wb:
        assert not ws.sheet_view.showGridLines
        for row in ws:
            for cell in row:
                if cell.data_type=="f":
                    assert "#REF!" not in cell.value
    assert wb["Dashboard"]._charts
    files=powerbi_files(b)
    assert len(files)==9
    periods=pd.read_csv(BytesIO(files["Periods.csv"]))
    forecasts=pd.read_csv(BytesIO(files["Forecasts.csv"]))
    assert periods.PeriodID.is_unique
    assert set(forecasts.PeriodID)<=set(periods.PeriodID)
    assert set(forecasts.Scenario)=={"Bear","Base","Bull"}
    actuals=pd.read_csv(BytesIO(files["FinancialActuals.csv"]))
    assert actuals.loc[(actuals.Metric=="Revenue")&(actuals.PeriodID=="2025-12-31"),"Value"].iloc[0]==1.1e9


def test_company_aliases():
    assert resolve("Apple")=="AAPL"
    assert resolve("Tata Consultancy Services")=="TCS.NS"
    assert resolve("RELIANCE.NS")=="RELIANCE.NS"


def test_no_shares_no_infinity(sample_data):
    del sample_data.info["sharesOutstanding"]
    b=run_model(sample_data)
    assert np.isnan(b.valuations["Base"].values["Implied Price"])
    assert b.forecasts["Base"].income.loc["EPS"].isna().all()
    excel_workbook(b,Units("USD"))


def test_snapshot_fallback(sample_data,tmp_path,monkeypatch):
    import finmodel_agents.market_data as module
    monkeypatch.setattr(module,"ROOT",tmp_path)
    agent=MarketDataAgent()
    agent._save(sample_data)
    loaded=agent._load("TEST")
    assert loaded.retrieved_at==sample_data.retrieved_at
    assert loaded.income.shape==sample_data.income.shape
