from io import BytesIO
from pathlib import Path
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from finmodel_agents.export import excel_workbook, excel_number
from finmodel_agents.orchestrator import run_model
from finmodel_agents.market_data import MarketDataAgent
from finmodel_agents.utils import Units


@pytest.mark.parametrize("missing",[None,float("nan"),float("inf"),float("-inf"),pd.NA])
def test_missing_valuation_cells_export_without_crash(sample_data,missing):
    bundle=run_model(sample_data)
    for valuation in bundle.valuations.values():
        for key in valuation.values:
            valuation.values[key]=missing
        valuation.sensitivity=pd.DataFrame(index=range(5),columns=range(5),dtype=float)
    raw=excel_workbook(bundle,Units("USD"))
    book=load_workbook(BytesIO(raw),data_only=True)
    assert book["DCF"]["B7"].value=="N/A"
    assert book["DCF"]["B11"].value=="N/A"
    for sheet in book:
        for row in sheet:
            for cell in row:
                assert cell.data_type!="e"
                if isinstance(cell.value,float):
                    assert np.isfinite(cell.value)


def test_zero_is_preserved_and_invalid_numbers_are_not_fabricated():
    assert excel_number(0)==0
    assert excel_number("12.5")==12.5
    assert excel_number(-1e7,1e7)==-1
    assert excel_number(10**1000)=="N/A"
    assert excel_number("unavailable")=="N/A"


@pytest.mark.parametrize("key",["marketCap","sharesOutstanding","currentPrice"])
def test_missing_market_input_still_exports(sample_data,key):
    sample_data.info[key]=None
    bundle=run_model(sample_data)
    book=load_workbook(BytesIO(excel_workbook(bundle,Units("USD"))),data_only=True)
    assert len(book.sheetnames)==21


def test_export_failure_does_not_break_dashboard(sample_data,monkeypatch):
    import app
    def fail(*args,**kwargs):
        raise TypeError("NAN/INF not supported in write_number()")
    monkeypatch.setattr(app,"export_excel",fail)
    at=AppTest.from_file(str(Path(__file__).resolve().parents[1]/"streamlit_app.py"),default_timeout=60)
    at.session_state["company"]=sample_data
    at.session_state["peers"]=[]
    at.run()
    assert not at.error
    assert not at.exception
    assert len(at.tabs)==10
    assert len(at.get("download_button"))==9
    assert any("Excel download is temporarily unavailable" in item.value for item in at.warning)
    assert any("Critical investment decisions" in item.value for item in at.caption)


def test_tcs_snapshot_without_market_cap_exports():
    data=MarketDataAgent()._load("TCS.NS")
    if data is None:
        pytest.skip("Run the live smoke test to obtain a TCS snapshot")
    data.info["marketCap"]=None
    bundle=run_model(data)
    assert not bundle.valuations["Base"].values
    raw=excel_workbook(bundle,Units("INR"))
    book=load_workbook(BytesIO(raw),data_only=True)
    assert "unavailable" in book["DCF"]["A6"].value.lower()
